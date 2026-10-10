"""Turns the rating sheet from `judge.py export-human` into one Excel file per rater.

CSV files with UTF-8 text and multi-line cells often open garbled in Excel, and free-text cells
invite answers such as "ya" or "4,5". Each workbook keeps the same columns, wraps the text,
freezes the header and restricts the rating cells to 0/1 or 1-5.

    python revision/human_sheet_xlsx.py --sheet revision/judge/human_sheet.csv --raters A B C
    # -> revision/judge/rater_A.xlsx, rater_B.xlsx, rater_C.xlsx

`judge.py agreement --human rater_A.xlsx rater_B.xlsx rater_C.xlsx ...` reads the returned files.
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

WIDTHS = {"item": 6, "article": 80, "assistant_output": 26, "retrieved_examples": 60, "comment": 30}
BINARY = ("_ok",)
SCALE = ("editorial_usefulness", "examples_helpful")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--raters", nargs="+", default=["A", "B", "C"])
    a = ap.parse_args()
    with open(a.sheet, encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header, body = rows[0], rows[1:]
    for r in a.raters:
        wb = Workbook()
        ws = wb.active
        ws.title = "penilaian"
        ws.append(header)
        for row in body:
            ws.append([int(row[0])] + row[1:])
        fill = PatternFill("solid", fgColor="FFF4CC")
        bin_dv = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=True,
                                error="Isi 1 (setuju) atau 0 (tidak setuju).", showErrorMessage=True)
        scale_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                                  error="Isi angka bulat 1 sampai 5.", showErrorMessage=True)
        ws.add_data_validation(bin_dv)
        ws.add_data_validation(scale_dv)
        last = len(body) + 1
        for j, name in enumerate(header, 1):
            col = ws.cell(row=1, column=j).column_letter
            ws.cell(row=1, column=j).font = Font(bold=True)
            ws.column_dimensions[col].width = WIDTHS.get(name, 14)
            rng = f"{col}2:{col}{last}"
            if name.endswith(BINARY):
                bin_dv.add(rng)
            elif name in SCALE:
                scale_dv.add(rng)
            if name.endswith(BINARY) or name in SCALE or name == "comment":
                for i in range(2, last + 1):
                    ws.cell(row=i, column=j).fill = fill
        for row in ws.iter_rows(min_row=1, max_row=last):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        out = Path(a.sheet).with_name(f"rater_{r}.xlsx")
        wb.save(out)
        print(f"✅ {out}")


if __name__ == "__main__":
    main()
